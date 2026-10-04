"""Untrusted document extraction. Production callers must use the bounded worker.

Nothing in this module admits facts, executes document instructions, opens linked
resources or writes attachments. Locators refer to the original submitted bytes.
"""
import csv
import hashlib
import io
import math
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Literal
from zipfile import ZipFile

from pydantic import BaseModel, ConfigDict, Field

MAX_INPUT = 5 * 1024 * 1024
MAX_EXPANDED = 20 * 1024 * 1024
MAX_STREAM = 4 * 1024 * 1024
MAX_TEXT = 200_000
MAX_CELLS = 10_000
MAX_ROWS = 2_000
MAX_COLUMNS = 100
MAX_PAGES = 100
MAX_SHEETS = 20


class ImportFailure(ValueError):
    """Only application-owned fixed error codes may cross the worker boundary."""


class ImportCell(BaseModel):
    model_config = ConfigDict(extra="forbid")
    locator: str = Field(max_length=32)
    text: str = Field(max_length=8192)
    kind: Literal["text", "number", "date", "boolean", "error", "formula"] = "text"
    number_format: str | None = Field(default=None, max_length=200)
    raw_value: str | None = Field(default=None, max_length=8192)


class ImportBlock(BaseModel):
    model_config = ConfigDict(extra="forbid")
    locator: str = Field(max_length=200)
    text: str = Field(default="", max_length=MAX_TEXT)
    cells: list[ImportCell] = Field(default_factory=list, max_length=MAX_COLUMNS)


class ParsedDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")
    format: Literal["csv", "xlsx", "pdf", "html"]
    content_hash: str = Field(pattern=r"^[a-f0-9]{64}$")
    verified: Literal[False] = False
    blocks: list[ImportBlock] = Field(max_length=MAX_ROWS)
    limitations: list[Literal["unverified_import", "formulas_not_evaluated", "pdf_layout_not_verified", "empty_pages", "merged_cells_not_expanded"]]


class Budget:
    def __init__(self):
        self.characters = self.cells = self.rows = 0

    def text(self, value):
        self.characters += len(value)
        if self.characters > MAX_TEXT or "\x00" in value:
            raise ImportFailure("content_limit")
        return value

    def row(self, cells):
        self.rows += 1
        self.cells += len(cells)
        if self.rows > MAX_ROWS or self.cells > MAX_CELLS or len(cells) > MAX_COLUMNS:
            raise ImportFailure("table_limit")


def text_encoding(raw):
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    return raw.decode(encoding, errors="strict")


def parse_csv(raw, budget):
    blocks = []
    reader = csv.reader(io.StringIO(text_encoding(raw), newline=""), strict=True)
    for index, row in enumerate(reader, 1):
        budget.row(row)
        cells = [ImportCell(locator=f"R{index}C{column}", text=budget.text(value)) for column, value in enumerate(row, 1)]
        blocks.append(ImportBlock(locator=f"row {index}", cells=cells))
    return blocks, ["formulas_not_evaluated"]


def validate_workbook_archive(raw):
    from defusedxml.ElementTree import fromstring
    from openpyxl.utils.cell import coordinate_to_tuple
    values = {}
    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 500 or len({entry.filename for entry in entries}) != len(entries):
            raise ImportFailure("archive_limit")
        if sum(entry.file_size for entry in entries) > MAX_EXPANDED:
            raise ImportFailure("archive_limit")
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if (entry.flag_bits & 1 or entry.file_size > MAX_STREAM or entry.file_size > max(1, entry.compress_size) * 200
                    or path.is_absolute() or ".." in path.parts or "\\" in entry.filename or ":" in entry.filename):
                raise ImportFailure("archive_limit")
            if ((path.suffix.lower() not in (".xml", ".rels") and path.name != ".rels")
                    or "vba" in entry.filename.casefold() or "externallinks" in entry.filename.casefold()):
                raise ImportFailure("workbook_active_content")
            content = archive.read(entry)
            # Preflight every XML part, including those a reader might otherwise ignore.
            root = fromstring(content, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            if any(element.attrib.get("TargetMode", "").casefold() == "external" for element in root.iter()):
                raise ImportFailure("workbook_external_reference")
            if any("macroenabled" in str(value).casefold() for element in root.iter() for value in element.attrib.values()):
                raise ImportFailure("workbook_active_content")
            if entry.filename.startswith("xl/worksheets/") and path.suffix == ".xml":
                sheet_values = {}
                for cell in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                    coordinate = cell.attrib.get("r", "")
                    row, column = coordinate_to_tuple(coordinate)
                    if row > MAX_ROWS or column > MAX_COLUMNS or coordinate in sheet_values:
                        raise ImportFailure("table_limit")
                    value = cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                    sheet_values[coordinate] = value.text if value is not None else None
                values[entry.filename] = sheet_values
    return values


def parse_xlsx(raw, budget):
    original_values = validate_workbook_archive(raw)
    import openpyxl
    from openpyxl.xml import DEFUSEDXML
    if not DEFUSEDXML:
        raise ImportFailure("safe_xml_unavailable")
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    blocks, formulas = [], False
    try:
        if len(book.worksheets) > MAX_SHEETS:
            raise ImportFailure("sheet_limit")
        for sheet in book.worksheets:
            # This pinned reader exposes the original part path; fail if it ever drifts.
            values = original_values.get(sheet._worksheet_path)
            if values is None:
                raise ImportFailure("unsupported_workbook_layout")
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ImportFailure("table_limit")
            # Ignore potentially understated dimension metadata; count actual cells/rows.
            sheet.reset_dimensions()
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                budget.row(row)
                cells = []
                for column, cell in enumerate(row, 1):
                    value, kind = cell.value, "text"
                    if value is None:
                        text = ""
                    elif cell.data_type == "f":
                        if not isinstance(value, str):
                            raise ImportFailure("unsupported_formula")
                        text, kind, formulas = value, "formula", True
                    elif cell.data_type == "e":
                        text, kind = str(value), "error"
                    elif isinstance(value, bool):
                        text, kind = str(value).lower(), "boolean"
                    elif isinstance(value, (datetime, date, time)):
                        text, kind = value.isoformat(), "date"
                    elif isinstance(value, (int, float)):
                        if not math.isfinite(value):
                            raise ImportFailure("invalid_number")
                        text, kind = str(value), "number"
                    else:
                        text = str(value)
                    original = values.get(getattr(cell, "coordinate", "")) if kind in ("number", "date", "formula") else None
                    if kind == "number" and original is not None:
                        text = original  # Never round an original decimal through Python float.
                    if original is not None:
                        budget.text(original)
                    cells.append(ImportCell(locator=f"R{row_index}C{column}", text=budget.text(text), kind=kind,
                                            number_format=getattr(cell, "number_format", None), raw_value=original))
                blocks.append(ImportBlock(locator=f"sheet {sheet.title} row {row_index}", cells=cells))
        return blocks, ["merged_cells_not_expanded", *(["formulas_not_evaluated"] if formulas else [])]
    finally:
        book.close()


def check_pdf_objects(reader):
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject
    pending, seen, count = [reader.trailer], set(), 0
    while pending:
        value = pending.pop()
        count += 1
        if count > 50_000 or len(pending) > 50_000:
            raise ImportFailure("pdf_object_limit")
        if isinstance(value, IndirectObject):
            key = (value.idnum, value.generation)
            if key in seen:
                continue
            seen.add(key)
            value = value.get_object()
        if isinstance(value, DictionaryObject):
            if any(key in value for key in ("/JS", "/JavaScript", "/AA", "/OpenAction", "/Launch", "/EmbeddedFiles", "/RichMedia", "/XFA")):
                raise ImportFailure("pdf_active_content")
            if value.get("/S") in ("/JavaScript", "/Launch", "/SubmitForm", "/ImportData", "/Rendition", "/GoToE"):
                raise ImportFailure("pdf_active_content")
            pending.extend(value.values())
        elif isinstance(value, ArrayObject):
            pending.extend(value)


def parse_pdf(raw, budget):
    from pypdf import PdfReader, apply_configuration
    if not raw.startswith(b"%PDF-"):
        raise ImportFailure("invalid_pdf")
    # No external decoder binary; compressed streams and recursive page/forms are bounded.
    with apply_configuration(maximum_declared_stream_length=MAX_STREAM, array_based_stream_maximum_output_length=MAX_STREAM,
            jbig2_maximum_output_length=MAX_STREAM, lzw_maximum_output_length=MAX_STREAM,
            run_length_maximum_output_length=MAX_STREAM, zlib_maximum_output_length=MAX_STREAM,
            image_maximum_buffer_size=MAX_STREAM, jbig2dec_binary=None,
            page_tree_maximum_entries=MAX_PAGES * 2, page_tree_maximum_depth=20,
            xform_maximum_invocations_per_extraction=100):
        reader = PdfReader(io.BytesIO(raw), strict=True, root_object_recovery_limit=1000)
        if reader.is_encrypted:
            raise ImportFailure("encrypted_pdf")
        check_pdf_objects(reader)
        if len(reader.pages) > MAX_PAGES:
            raise ImportFailure("page_limit")
        blocks, empty = [], False
        for index, page in enumerate(reader.pages, 1):
            content = page.get_contents()
            if content and len(content.get_data()) > MAX_STREAM:
                raise ImportFailure("pdf_stream_limit")
            text = budget.text(page.extract_text())
            empty = empty or not text.strip()
            blocks.append(ImportBlock(locator=f"page {index}", text=text))
        if not any(block.text.strip() for block in blocks):
            raise ImportFailure("pdf_text_unavailable")
        return blocks, ["pdf_layout_not_verified", *(["empty_pages"] if empty else [])]


def parse_document(raw: bytes, format: str, *, permission_confirmed: bool) -> ParsedDocument:
    """Worker-only synchronous parser. Permission is separate from factual verification."""
    if not permission_confirmed:
        raise ImportFailure("permission_required")
    if not raw or len(raw) > MAX_INPUT:
        raise ImportFailure("input_limit")
    budget = Budget()
    try:
        if format == "html":
            from backend.app.evidence import VisibleText
            parser = VisibleText()
            parser.feed(text_encoding(raw))
            blocks, limitations = [ImportBlock(locator="document", text=budget.text(parser.text()))], []
        else:
            parser = {"csv": parse_csv, "xlsx": parse_xlsx, "pdf": parse_pdf}.get(format)
            if parser is None:
                raise ImportFailure("unsupported_format")
            blocks, limitations = parser(raw, budget)
        if not blocks:
            raise ImportFailure("empty_document")
        return ParsedDocument(format=format, content_hash=hashlib.sha256(raw).hexdigest(), blocks=blocks,
                              limitations=["unverified_import", *limitations])
    except ImportFailure:
        raise
    except Exception:
        raise ImportFailure("malformed_or_unsupported_document") from None
